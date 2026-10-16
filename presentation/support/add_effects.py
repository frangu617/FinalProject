from pptx import Presentation
from pptx.oxml.xmlchemy import OxmlElement
from pathlib import Path
ROOT = Path(__file__).resolve().parents[2]
p=Presentation(ROOT/'presentation/originals/before-effects/poker_hand_presentation.pptx')
def el(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items(): e.set(k,str(v))
 return e
def child(parent,tag,**attrs):
 e=el(tag,**attrs);parent.append(e);return e

def timing(slide,steps):
 timing=el('p:timing');tn=child(timing,'p:tnLst');par=child(tn,'p:par');root=child(par,'p:cTn',id=1,dur='indefinite',restart='never',nodeType='tmRoot');cl=child(root,'p:childTnLst');seq=child(cl,'p:seq',concurrent=1,nextAc='seek');main=child(seq,'p:cTn',id=2,dur='indefinite',nodeType='mainSeq');maincl=child(main,'p:childTnLst');ident=3
 for step in steps:
  par=child(maincl,'p:par');ct=child(par,'p:cTn',id=ident,fill='hold');ident+=1
  child(child(ct,'p:stCondLst'),'p:cond',delay='indefinite')
  inner=child(child(ct,'p:childTnLst'),'p:par');ct2=child(inner,'p:cTn',id=ident,fill='hold');ident+=1
  child(child(ct2,'p:stCondLst'),'p:cond',delay=0)
  effects=child(ct2,'p:childTnLst')
  for effect_index,(shape_id,paragraph) in enumerate(step):
   pp=child(effects,'p:par');effect=child(pp,'p:cTn',id=ident,presetID=10,presetClass='entr',presetSubtype=0,fill='hold',grpId=0,nodeType='clickEffect' if effect_index==0 else 'withEffect');ident+=1
   child(child(effect,'p:stCondLst'),'p:cond',delay=0)
   ec=child(effect,'p:childTnLst')
   def behavior(parent,duration):
    return child(parent,'p:cBhvr')
   se=child(ec,'p:set');b=behavior(se,1);child(b,'p:cTn',id=ident,dur=1,fill='hold');ident+=1
   target=child(child(b,'p:tgtEl'),'p:spTgt',spid=shape_id)
   if paragraph is not None: child(child(target,'p:txEl'),'p:pRg',st=paragraph,end=paragraph)
   child(child(b,'p:attrNameLst'),'p:attrName').text='style.visibility';child(child(se,'p:to'),'p:strVal',val='visible')
   ae=child(ec,'p:animEffect',transition='in',filter='fade');b=behavior(ae,350);child(b,'p:cTn',id=ident,dur=350);ident+=1
   target=child(child(b,'p:tgtEl'),'p:spTgt',spid=shape_id)
   if paragraph is not None:child(child(target,'p:txEl'),'p:pRg',st=paragraph,end=paragraph)
 for name,event in [('p:prevCondLst','onPrev'),('p:nextCondLst','onNext')]:
  cond=child(child(seq,name),'p:cond',evt=event,delay=0);child(child(cond,'p:tgtEl'),'p:sldTgt')
 builds=child(timing,'p:bldLst')
 for sid in sorted({sid for step in steps for sid,pa in step}):
  pars=any(pa is not None for step in steps for shape,pa in step if shape==sid)
  child(builds,'p:bldP',spid=sid,grpId=0,build='p' if pars else 'allAtOnce')
 slide._element.append(timing)

import json
plan=json.loads((ROOT/'presentation/support/slide_plan.json').read_text())
# One sentence group per click, aligned to the visible paragraph or table row.
spoken={
3:[
"Our goal is to classify a supplied five-card hand into one of ten categories. We are not predicting future cards or an opponent's hidden cards.",
"Our experience with statistics has often centered on medical research. Poker gives us a different setting in which to apply the same ideas.",
"We chose the UCI dataset because its known labels and mix of common and rare hands let us study class imbalance. Our main question is how performance differs between these categories. Because rules determine the correct category, we can also ask what a statistical model actually learns."],
5:[
"The supplied files contain twenty-five thousand and ten training records and one million test records. Training records pass checks for missing values, valid codes, and repeated cards within a hand.",
"We identify the same five cards even when their order changes. Removing one hundred and twelve repeated hands leaves twenty-four thousand eight hundred and ninety-eight distinct development hands.",
"We split these into nineteen thousand nine hundred and eighteen training hands and four thousand nine hundred and eighty validation hands. The supplied test file stays separate from fitting and model selection.",
"We one-hot encode suits and ranks so their codes are treated as discrete values. The encoder learns only from the training split."],
7:[
"The main exploratory finding is severe class imbalance. High card represents about fifty percent of the combined files, and one pair represents about forty-two percent.",
"Together, they account for more than ninety-two percent of records.",
"There are only eight royal-flush records. These are descriptive frequencies across both supplied files, not held-out model results. We disclose that the test file was included in that earlier descriptive exploration.",
"A model can score reasonably well on accuracy while missing uncommon hands. That is why we need metrics that examine all ten categories."],
9:[
"We compare three candidates on the same validation split. Our reference is a majority baseline that always predicts high card.",
"The decision tree is a depth-limited learned comparison.",
"The random forest combines two hundred trees. Class weighting increases the influence of uncommon categories but does not create examples. Both learned models use one-hot encoding fitted on training data.",
"We select using macro F1, which averages F1 across all ten categories, and also report accuracy and class-level results. Lilian will now explain the implementation and findings."],
11:[
"The random forest has the highest validation macro F1 among these candidates, so the notebook selects it.",
"Its validation accuracy is sixty point five six percent.",
"However, macro F1 is only zero point one one nine four. That is a warning that performance is uneven across categories.",
"Training accuracy is one hundred percent. The large training-to-validation gap suggests overfitting. Selection means the strongest of these three candidates on this one split, not an established reliable classifier across all categories."],
12:[
"The majority baseline reaches forty-nine point nine six percent accuracy, with macro F1 of zero point zero six six six.",
"The decision tree has lower accuracy, twenty-two point five seven percent, but higher macro F1 of zero point zero eight three seven. This shows why the two metrics can rank models differently.",
"The forest leads on both measures: sixty point five six percent accuracy and macro F1 of zero point one one nine four.",
"Its accuracy exceeds the baseline by ten point six percentage points. All these numbers come from the four thousand nine hundred and eighty validation hands. The optional final test evaluation has not been run."],
14:[
"The class-level results explain the weak macro F1. The forest correctly identifies about eighty-eight percent of actual high-card hands.",
"It identifies only about forty percent of actual one-pair hands. Even the notebook's example pair of tens is predicted as high card.",
"The other eight categories have zero recall on this validation split.",
"A single example illustrates an error, but the class-level evaluation is the broader evidence. We also need support, the number of examples behind each score."],
16:[
"Each of the three rarest categories has only one validation example. This cannot establish reliable population performance.",
"The comparison uses one validation split, which also guides selection. We have not measured variability across repeated evaluations.",
"Inputs retain card positions, so changing card order may change a prediction. Deduplication removes identical card sets but does not group suit-renamed equivalents.",
"Final testing and development-to-test overlap reporting remain pending. We also disclose the earlier descriptive exploration of both supplied files."],
17:[
"Our next experiments should improve representation. We can compare consistent card ordering with order-independent features.",
"A correct rule-based evaluator provides a useful reference. We should explain when engineered features encode poker rules directly.",
"We should repeat validation where rare-class counts allow it.",
"Then we should lock the model and procedure before final testing. Report results on the supplied test file and separately on hands absent from development. These are proposed next steps, not completed results."],
18:[
"Our answer is qualified: the random forest improves validation accuracy over a majority baseline.",
"Low macro F1, zero recall outside the two most common classes, and the training-to-validation gap limit that result.",
"The current representation does not support reliable classification across all ten categories. The project shows why careful preparation and class-level interpretation matter as much as one headline accuracy."],
}
paragraph_slides={3:170,5:184,7:198,9:211,11:225,14:255,16:268,17:274,18:280}
picture_slides={4:179,6:193,8:205,10:220,13:249,15:262}
labels={4:'the five-card record example',6:'the preparation code',8:'the class-distribution chart',10:'the model pipeline code',13:'the two-metric comparison chart',15:'the class-recall chart'}
manifest=[]
notes_md=['# Two-speaker presentation notes with animation cues','',
'Target: approximately 9–10 minutes including short pauses. Francisco presents slides 1–9; Lilian presents slides 10–18; both share slide 19. Rehearse with the builds enabled.', '',
'Controls: use a single click, Space, or Right Arrow for each numbered reveal. After the last reveal, the NEXT SLIDE click changes slides. Wait about half a second for each fade/dissolve; avoid double-clicks. Slide changes are manual. The PDF is static and shows all content.', '']
for n,(slide,item) in enumerate(zip(p.slides,plan),1):
 # Keynote's PowerPoint importer drops zero-opacity text; remove the hidden code duplicate.
 hidden_ids={4:177,6:191,10:218}
 if n in hidden_ids:
  duplicate=next(shape for shape in slide.shapes if shape.shape_id==hidden_ids[n])
  duplicate._element.getparent().remove(duplicate._element)
 for tag in ['p:timing','p:transition']:
  for existing in list(slide._element.findall(tag,slide._element.nsmap)): slide._element.remove(existing)
 if n<19:
  transition=el('p:transition',spd='fast',advClick=1);child(transition,'p:fade');slide._element.append(transition)
 steps=[];names=[]
 if n in paragraph_slides:
  sid=paragraph_slides[n];shape=next(s for s in slide.shapes if s.shape_id==sid)
  assert len(shape.text_frame.paragraphs)==len(spoken[n])
  for i,paragraph in enumerate(shape.text_frame.paragraphs):
   steps.append([(sid,i)]);names.append(paragraph.text)
 elif n in picture_slides:
  steps=[[(picture_slides[n],None)]];names=[labels[n]]
 elif n==12:
  steps=[[(sid,None) for sid in [234,235,236]],[(sid,None) for sid in [237,238,239]],[(sid,None) for sid in [240,241,242]],[(243,None)]]
  names=['the majority-baseline row','the decision-tree row','the random-forest row','the +10.60 percentage-point takeaway']
 if steps: timing(slide,steps)
 lines=[f'SPEAKER: {item["speaker"]}',f'REHEARSAL TARGET: {item["seconds"]} seconds plus brief reveal pauses',f'REVEALS: {len(steps)}; then one separate click to advance.', '', '[ON ARRIVAL — let the 0.45-second dissolve finish; pause 0.5 seconds. Title and context are visible. Do not advance again yet.]','']
 if n in spoken:
  for j,(name,words) in enumerate(zip(names,spoken[n]),1):
   lines += [f'[CLICK {j} — fade in {name} Wait 0.5 seconds, then speak.]',words,'']
 elif steps:
  lines += [f'[CLICK 1 — fade in {names[0]}. Wait 0.5 seconds'+('; allow 2 seconds for the audience to inspect the chart.]' if n in [8,13,15] else '; pause 1 second before explaining.]'),item['script'],'']
 else:
  lines += ['[NO CONTENT BUILDS — keep this slide visible while speaking.]',item['script'],'']
 if n==9:
  lines += ['[HANDOFF — look toward Lilian and pause 1 second. Advance once; Lilian triggers the next slide\'s code reveal.]','']
 elif n==18:
  lines += ['[PAUSE 1 second after the conclusion. Advance once to the shared closing slide.]','']
 elif n==19:
  lines += ['[FINISH — pause 2 seconds after “Thank you.” Leave the closing slide visible; do not advance.]','']
 else:
  lines += [f'[NEXT SLIDE — after the final sentence'+(' and all '+str(len(steps))+' reveals' if steps else '')+', pause 1 second. Click once to advance. Let the next dissolve finish before triggering a reveal.]','']
 lines += ['EVIDENCE: '+item['source']]
 note='\n'.join(lines);slide.notes_slide.notes_text_frame.text=note
 notes_md += [f'## {n}. {item["title"]}','',note,'']
 manifest.append({'slide':n,'title':item['title'],'speaker':item['speaker'],'reveal_clicks':len(steps),'groups':steps,'labels':names,'notes':note})
(ROOT/'presentation/two_speaker_notes.md').write_text('\n'.join(notes_md))
(ROOT/'presentation/support/effects_manifest.json').write_text(json.dumps(manifest,indent=2))
p.save(ROOT/'presentation/support/animated_candidate.pptx')
print('Created 19-slide candidate with',sum(x['reveal_clicks'] for x in manifest),'click-to-reveal groups.')
